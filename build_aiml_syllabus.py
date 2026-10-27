from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "AIML_Syllabus.html"
OUTPUT = ROOT / "AIML_Syllabus.docx"
PART_COLORS = ["237A57", "286D9C", "B66A24", "167C79", "B34F45", "536574"]
PALE_COLORS = ["EAF5EE", "EAF3F9", "FFF3E5", "E7F5F3", "FBEFED", "EFF2F4"]


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def borders(cell, color="D7E0E5", size="6", left_color=None):
    properties = cell._tc.get_or_add_tcPr()
    border_set = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "18" if edge == "left" and left_color else size)
        border.set(
            qn("w:color"), left_color if edge == "left" and left_color else color
        )
        border_set.append(border)
    properties.append(border_set)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def add_bullets(cell, items):
    for item in items:
        paragraph = cell.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Inches(0.13)
        paragraph.paragraph_format.first_line_indent = Inches(-0.1)
        run = paragraph.add_run(item)
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(52, 74, 89)


def fill_card(cell, title, items, color):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell)
    borders(cell, left_color=color)
    heading = cell.paragraphs[0]
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    add_bullets(cell, items)


def add_card_rows(document, cards, color):
    for offset in range(0, len(cards), 2):
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        for column, width in enumerate((3.55, 3.55)):
            table.columns[column].width = Inches(width)
            table.cell(0, column).width = Inches(width)
        for column in range(2):
            index = offset + column
            cell = table.cell(0, column)
            if index < len(cards):
                title, items = cards[index]
                fill_card(cell, title, items, color)
            else:
                cell.text = ""
                borders(cell, color="FFFFFF", size="0")
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1)
        spacer.paragraph_format.space_before = Pt(0)


def card_content(article):
    heading = article.find("h3")
    items = [item.get_text(" ", strip=True) for item in article.select("li")]
    return heading.get_text(" ", strip=True), items


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite {OUTPUT}")

    soup = BeautifulSoup(SOURCE.read_text(encoding="utf-8"), "html.parser")
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(23, 42, 58)
    normal.paragraph_format.space_after = Pt(4)

    banner = document.add_table(rows=1, cols=1)
    banner.autofit = False
    banner.columns[0].width = Inches(7.1)
    banner_cell = banner.cell(0, 0)
    shade(banner_cell, "174E4D")
    set_cell_margins(banner_cell, top=260, start=240, bottom=220, end=240)
    eyebrow = banner_cell.paragraphs[0]
    eyebrow.paragraph_format.space_after = Pt(3)
    run = eyebrow.add_run("COMPLETE LEARNING ROADMAP")
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(215, 239, 223)
    title = banner_cell.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("AI & Machine Learning")
    run.bold = True
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor(255, 255, 255)
    subtitle = banner_cell.add_paragraph()
    run = subtitle.add_run(
        "A chapter-wise syllabus from Python foundations to production-ready generative AI"
    )
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(242, 248, 242)

    intro = soup.select_one(".intro")
    paragraph = document.add_paragraph(intro.get_text(" ", strip=True))
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(83, 101, 116)

    facts = soup.select(".quick-facts .fact")
    fact_table = document.add_table(rows=1, cols=len(facts))
    fact_table.autofit = False
    for index, fact in enumerate(facts):
        cell = fact_table.cell(0, index)
        cell.width = Inches(7.1 / max(1, len(facts)))
        shade(cell, "F3F7F3")
        borders(cell, color="D7E0E5", left_color="237A57")
        set_cell_margins(cell, top=90, start=120, bottom=90, end=120)
        strong = fact.find("strong")
        first = cell.paragraphs[0]
        first.paragraph_format.space_after = Pt(1)
        run = first.add_run(strong.get_text(" ", strip=True))
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("237A57")
        detail = cell.add_paragraph(fact.find("span").get_text(" ", strip=True))
        detail.paragraph_format.space_after = Pt(0)
        for run in detail.runs:
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor(83, 101, 116)

    for part_index, part in enumerate(soup.select(".part")):
        color = PART_COLORS[part_index]
        pale = PALE_COLORS[part_index]
        name = part.select_one(".part-heading h2").get_text(" ", strip=True)
        heading_table = document.add_table(rows=1, cols=1)
        heading_cell = heading_table.cell(0, 0)
        shade(heading_cell, pale)
        borders(heading_cell, color=pale, left_color=color)
        set_cell_margins(heading_cell, top=80, start=140, bottom=80, end=140)
        heading_paragraph = heading_cell.paragraphs[0]
        heading_paragraph.paragraph_format.space_after = Pt(0)
        run = heading_paragraph.add_run(f"PART {part_index + 1:02d}  |  {name}")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string(color)

        chapter_cards = [
            card_content(article)
            for article in part.select(".chapter-grid article.chapter")
        ]
        if chapter_cards:
            add_card_rows(document, chapter_cards, color)

        for article in part.select(
            ":scope > article.capstone, :scope > article.specializations"
        ):
            title_text, items = card_content(article)
            table = document.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            if "specializations" in article.get("class", []):
                shade(cell, "FFF9EC")
                color_for_card = "B66A24"
            else:
                shade(cell, "EAF5EE")
                color_for_card = "237A57"
            fill_card(cell, title_text, items, color_for_card)
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(1)
            spacer.paragraph_format.space_before = Pt(0)

        tools = part.select(".tools .tool")
        if tools:
            label = document.add_paragraph()
            label.paragraph_format.space_before = Pt(4)
            label.paragraph_format.space_after = Pt(4)
            run = label.add_run("Suggested Practical Tools")
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor.from_string(color)
            tool_cards = []
            for tool in tools:
                strong = tool.find("strong")
                lead = strong.get_text(" ", strip=True)
                remainder = tool.get_text(" ", strip=True)[len(lead) :].strip()
                tool_cards.append((lead, [remainder]))
            add_card_rows(document, tool_cards, color)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("AI & Machine Learning  |  Complete chapter-wise syllabus")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(83, 101, 116)

    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
